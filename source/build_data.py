# 把三份 Excel 整理成網頁用的 data.json（只有資料換了才需要重跑；平常改版面不用）
#
#   python3 build_data.py <新聞編碼.xlsx> <文獻.xlsx> <學生人數.xlsx> data.json
#
#   新聞編碼.xlsx   要有「2026整理」與「0815合併原始檔」兩張工作表
#   文獻.xlsx       要有「Design_themes(n=426)」與「Terms」兩張工作表
#   學生人數.xlsx   要有「通報網全台數據」「台北市年報＿各障別」「台北市年報＿特教學校」「衛福部全台數據」四張工作表
#
# 需要先安裝 openpyxl（pip install openpyxl）。
# 這支程式只讀取上面列出的欄位；「2026整理」的備註欄（有當事人姓名）不會讀進來。
# 跑完會印出一份檢查報告（篇數、對不上的日期、各種編碼的次數），請對一下有沒有和預期不同的地方。
# 原始的 Excel 檔含有個人資料，請不要放進公開的 GitHub 倉庫。
import sys, re, json, collections, datetime, os
import openpyxl
if len(sys.argv) != 5:
    sys.exit('用法：python3 build_data.py <新聞編碼.xlsx> <文獻.xlsx> <學生人數.xlsx> <輸出的 data.json>')
NEWS_XLSX, LIT_XLSX, COUNT_XLSX, OUT = sys.argv[1:5]
C = collections.Counter
report = []

# ---------------- 新聞 ----------------
wb = openpyxl.load_workbook(NEWS_XLSX, data_only=True)
ws = wb['2026整理']; rows = list(ws.iter_rows(values_only=True))
D = [r for r in rows[2:] if r[0] is not None]
wsf = openpyxl.load_workbook(NEWS_XLSX)['2026整理']      # 再讀一次，這次讀儲存格的底色：黃底（FFF2CC）是「有出現關鍵字但主體非智能障礙者」的標記
flag = {}
for i in range(3, wsf.max_row+1):
    c = wsf[f'A{i}']
    if c.value is not None and c.fill is not None and c.fill.fill_type and c.fill.fgColor.type == 'rgb':
        flag[str(c.value).strip()] = c.fill.fgColor.rgb
raw = wb['0815合併原始檔']; rr = list(raw.iter_rows(values_only=True))
rawmap = {str(r[1]).strip(): r for r in rr[1:] if r[1]}

OUTLETS = ['聯合晚報', '中國時報', '中央日報', '經濟日報', '聯合報', '美聯社']
SEP = r'[_＿\s]'
def parse_name(name):
    s = re.sub(r'\.(png|pdf|jpg|jpeg)$', '', str(name).strip(), flags=re.I)
    m = re.match(r'^(\d{4})(\d{2})(\d{2})' + SEP + r'*(.*)$', s)
    y, mo, d, rest = m.groups()
    outlet = ''
    for o in OUTLETS:
        if rest.startswith(o): outlet = o; rest = rest[len(o):]; break
    rest = re.sub('^' + SEP + '+', '', rest)
    rest2 = re.sub(r'^(第?[A-Za-z]?\d{1,2}版?|[一二三四五六七八九十]+版(本報訊)?)', '', rest)     # 版面
    if rest2 != rest and (rest2 == '' or re.match(SEP, rest2) or not re.search(SEP, rest)): rest = rest2
    elif re.search(SEP, rest) and not re.search(SEP, rest2): rest = rest2
    title = re.sub('^' + SEP + '+', '', rest)
    title = re.sub(r'^第\d{1,2}版', '', title)
    title = re.sub(SEP + '+', ' ', title).strip() or '（沒有標題）'
    return f'{y}-{mo}-{d}', outlet, title
def sp(v, extra=''):
    if v is None: return []
    return [x.strip() for x in re.split(r'[,' + extra + r']\s*', str(v)) if x.strip()]

TERM_GROUPS = [   # 圓圈圖與背景用的用詞分組；沒列到的歸「其他用詞」
    ('白痴／低能', ['白痴/低能', '低能', '白痴']),
    ('智能不足', ['智能不足']),
    ('智障', ['智障', '智障朋友']),
    ('弱智', ['弱智', '弱智者', '弱智兒']),
    ('憨兒／喜憨兒', ['憨兒', '喜憨兒', '憨孩子', '憨媽']),
    ('心智障礙', ['心智障礙', '心智障礙者', '心障']),
    ('智能障礙', ['智能障礙', '智能障礙青年']),
]
TG = {v: g for g, vs in TERM_GROUPS for v in vs}
VIEW_SHORT = {'沈重負擔/同情憐憫': '負擔・憐憫', '肯定障礙者(學習能力、成就)': '肯定能力', '倡議': '倡議',
    '家人/專業人員犧牲奉獻': '犧牲奉獻', '期待正常化/健全主義': '期待正常化', '排斥障礙者': '排斥', '慈善': '慈善',
    '優生保健': '優生保健', '家長拒絕標籤': '拒絕標籤', '勵志激勵': '勵志', '從抗拒到接受': '抗拒到接受'}
def topic_group(t):
    if t.startswith('專家學者'): return '專家學者'
    if t.startswith('表揚'): return '表揚'
    return t.split('-')[0]
CATS = ['事件報導', '經驗抒發、觀點與評論', '資訊分享']
def cat_norm(c):
    if c.startswith('資訊分享'): return '資訊分享'
    return c if c in CATS else '其他'
SUBJ = ['智能障礙者', '母親', '父親', '手足', '延伸家庭（祖父母、姻親、其他親戚）', '教師', '醫療專業人員', '同儕']

news = []; ymis = []; nojoin = 0; badtopic = C(); otherterms = C()
for r in D:
    name = str(r[0]).strip()
    date, outlet, title = parse_name(name)
    rw = rawmap.get(name)
    if rw and isinstance(rw[3], datetime.datetime): date = rw[3].strftime('%Y-%m-%d')
    else: nojoin += 1
    if rw and rw[4]: outlet = str(rw[4]).strip()
    if int(date[:4]) != int(r[1]): ymis.append((name, int(r[1]), date))
    terms = [re.sub(r'（.*?）', '', t).strip() for t in sp(r[2], '、')]
    terms = [t for t in terms if t]
    wg = []
    for t in terms:
        g = TG.get(t, '其他用詞')
        if g == '其他用詞': otherterms[t] += 1
        if g not in wg: wg.append(g)
    topics = [t for t in sp(r[6]) if not t.isdigit()]
    for t in sp(r[6]):
        if t.isdigit(): badtopic[t] += 1
    tg = []
    for t in topics:
        g = topic_group(t)
        if g not in tg: tg.append(g)
    v1 = r[7] if r[7] not in (None, '') else None
    v2 = r[8] if r[8] not in (None, '') else None
    cats = []
    for c in sp(r[5]):
        c = cat_norm(c)
        if c not in cats: cats.append(c)
    subj = []
    for s_ in sp(r[3]):
        s2 = s_ if s_ in SUBJ else '其他'
        if s2 not in subj: subj.append(s2)
    news.append({'d': date, 'o': outlet, 't': title, 'w': terms, 'wg': wg, 'p': topics, 'pg': tg,
                 'v': v1, 'v2': v2,
                 'x': 1 if flag.get(name) == 'FFFFF2CC' else 0})
news.sort(key=lambda a: a['d'])
report.append(f'新聞 {len(news)} 篇；{len(D)-nojoin} 篇對得到原始檔的日期，{nojoin} 篇用檔名的日期')
report.append(f'年份欄和日期不一致：{ymis}')
report.append(f'主題欄的雜訊值：{dict(badtopic)}')
report.append(f'標記「主體非智能障礙者」：{sum(a["x"] for a in news)} 篇')
report.append(f'歸到其他用詞：{sum(otherterms.values())} 次，{len(otherterms)} 種：{otherterms.most_common(60)}')
report.append(f'報別：{C(a["o"] for a in news).most_common()}')
report.append(f'觀點（主要）：{C(a["v"] for a in news).most_common()}')
report.append(f'主題大類：{C(g for a in news for g in a["pg"]).most_common()}')
report.append(f'用詞分組：{C(g for a in news for g in a["wg"]).most_common()}')
report.append(f'沒有主題的：{sum(1 for a in news if not a["p"])}；沒有用詞的：{sum(1 for a in news if not a["w"])}')

# ---------------- 文獻 ----------------
wb = openpyxl.load_workbook(LIT_XLSX, data_only=True)
ws = wb['Design_themes(n=426)']; rows = list(ws.iter_rows(values_only=True))
L = [r for r in rows[1:] if r[0] is not None]
DESIGNS = ['實驗或介入研究', '量化研究', '質性研究', '行動研究', '混合研究', '文獻回顧']
THEMES = [k for k, _ in C(r[3] for r in L).most_common()]
RIGHTS = ['性相關調查', '自我倡議/自我決策', '性教育介入']
lit = [[int(r[1]), DESIGNS.index(r[2]), THEMES.index(r[3]), (RIGHTS.index(r[4]) + 1) if r[4] in RIGHTS else 0, 1 if r[6] == '有' else 0] for r in L]
lit.sort()
ws = wb['Terms']; rows = list(ws.iter_rows(values_only=True))
def lit_term(t):
    t = str(t).strip()
    if t == '智能不足': return 0
    if t.startswith('智能障礙') or t.startswith('主要使用智能障礙'): return 1
    return 2
T = [(int(r[0]), lit_term(r[1]), str(r[1]).strip()) for r in rows[1:] if r[0] is not None and r[1] is not None]
litterms = sorted([y, c] for y, c, _ in T)
report.append(f'文獻研究 {len(lit)} 篇 {lit[0][0]}–{lit[-1][0]}；設計 {C(DESIGNS[a[1]] for a in lit).most_common()}')
report.append(f'文獻主題 {C(THEMES[a[2]] for a in lit).most_common()}')
report.append(f'文獻用詞 {len(litterms)} 筆：{C(c for _, c in litterms)}；其他={C(n for _, c, n in T if c == 2).most_common()}')

# ---------------- 人數 ----------------
wb = openpyxl.load_workbook(COUNT_XLSX, data_only=True)
def grid(ws): return [list(r) for r in ws.iter_rows(values_only=True)]
def series(g, yrow, drow, off=0, first=None):
    out = {}
    for j, y in enumerate(g[yrow]):
        if j == 0 or not isinstance(y, (int, float)): continue
        v = g[drow][j + off] if j + off < len(g[drow]) else None
        if isinstance(v, (int, float)): out[int(y)] = int(v)
    return out
g = grid(wb['通報網全台數據'])
nation = {k: series(g, 0, i) for i, k in enumerate(['ID', 'ASD', 'LD', 'CP', 'OTHER'], start=1)}
nation['TOTAL'] = series(g, 0, 6, off=1)
g = grid(wb['台北市年報＿各障別'])
taipei = {k: series(g, 1, i) for i, k in enumerate(['ID', 'ASD', 'LD', 'CP', 'OTHER'], start=2)}
taipei['TOTAL'] = series(g, 1, 7, off=1)
g2 = grid(wb['台北市年報＿特教學校'])
assert g2[33][0] == 'ID' and g2[32][0] == '四所特校加總'
taipei['SCHOOL_ID'] = series(g2, 1, 33)
g = grid(wb['衛福部全台數據'])
mohw = series(g, 1, 2)
for reg in (nation, taipei):                     # 腦性麻痺較早的年度填的是 0，改成沒有資料（圖上不畫）
    reg['CP'] = {y: v for y, v in reg['CP'].items() if v > 0}
counts = {'nation': nation, 'taipei': taipei, 'mohw': mohw}
for k, reg in (('全台', nation), ('台北', taipei)):
    for s_, d in reg.items(): report.append(f'人數 {k} {s_}: {min(d)}–{max(d)} {len(d)} 年，首 {d[min(d)]} 末 {d[max(d)]}')
report.append(f'衛福部: {min(mohw)}–{max(mohw)}，首 {mohw[min(mohw)]} 末 {mohw[max(mohw)]}')

json.dump({'news': news, 'viewShort': VIEW_SHORT, 'lit': lit, 'litDesigns': DESIGNS, 'litThemes': THEMES, 'litRights': RIGHTS,
           'litTerms': litterms, 'counts': counts}, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))
print('\n'.join(report)); print('bytes', os.path.getsize(OUT))
